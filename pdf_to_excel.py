# -*- coding: utf-8 -*-
"""
pdf_to_excel.py
================
แปลงไฟล์ "รายงานข้อมูลการใช้ไฟฟ้า" (PDF) ให้เป็นไฟล์ Excel (.xlsx)
สคริปต์นี้จัดการให้ทุกอย่างในตัวเอง ไม่ต้องติดตั้งอะไรเองล่วงหน้า

วิธีใช้ (เลือกแบบใดแบบหนึ่ง):
  1) ลากไฟล์ PDF (ไฟล์เดียวหรือหลายไฟล์) มาวางบนไอคอนของสคริปต์นี้
  2) วางไฟล์ PDF ไว้ในโฟลเดอร์เดียวกับสคริปต์นี้ แล้วดับเบิลคลิกที่สคริปต์ได้เลย
     (จะค้นหาและแปลงไฟล์ .pdf ทุกไฟล์ในโฟลเดอร์นั้นให้อัตโนมัติ)
  3) รันผ่าน Command Prompt:
         python pdf_to_excel.py "ไฟล์รายงาน.pdf"
         python pdf_to_excel.py "ไฟล์รายงาน.pdf" "ผลลัพธ์.xlsx"

ครั้งแรกที่รัน สคริปต์จะติดตั้งไลบรารีที่ต้องใช้ (pdfplumber, openpyxl) ให้อัตโนมัติ
(ต้องมี Python และอินเทอร์เน็ตสำหรับการติดตั้งครั้งแรกเท่านั้น)
"""

import sys
import os
import subprocess
import importlib

# ---------------------------------------------------------------------------
# 1) ติดตั้งไลบรารีที่ต้องใช้โดยอัตโนมัติ ถ้ายังไม่มี
# ---------------------------------------------------------------------------
REQUIRED_PACKAGES = ['pdfplumber', 'openpyxl']


def ensure_packages(packages):
    missing = []
    for pkg in packages:
        try:
            importlib.import_module(pkg)
        except ImportError:
            missing.append(pkg)
    if missing:
        print(f'กำลังติดตั้งไลบรารีที่ต้องใช้ ({", ".join(missing)}) กรุณารอสักครู่ ...')
        try:
            subprocess.check_call([sys.executable, '-m', 'pip', 'install', '--quiet', *missing])
            print('ติดตั้งไลบรารีเรียบร้อยแล้ว\n')
        except subprocess.CalledProcessError:
            print('ติดตั้งไลบรารีไม่สำเร็จ กรุณาเปิด Command Prompt แล้วพิมพ์คำสั่งนี้เอง:')
            print(f'    pip install {" ".join(missing)}')
            input('\nกด Enter เพื่อปิดหน้าต่างนี้...')
            sys.exit(1)


ensure_packages(REQUIRED_PACKAGES)

import re
import glob

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------------------------------------------------------------------------
# 2) ค่าคงที่สำหรับโครงสร้างตารางและการจัดรูปแบบ Excel
# ---------------------------------------------------------------------------
COLUMNS = ['หน่วยงาน', 'ชั้น', 'โซน', 'บล็อก', 'หมวด', 'Number',
           'On-Peak Unit', 'Off-Peak Unit', 'Total Unit']

FONT_NAME = 'TH Sarabun New'
FONT_SIZE = 14
HEADER_FILL = PatternFill('solid', start_color='1F4E78', end_color='1F4E78')
HEADER_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color='FFFFFF')
NORMAL_FONT = Font(name=FONT_NAME, size=FONT_SIZE)
THIN = Side(style='thin', color='B7C6D8')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUMFMT = '#,##0'


# ---------------------------------------------------------------------------
# 3) อ่านและแปลงข้อมูลจาก PDF
# ---------------------------------------------------------------------------
def extract_lines(pdf_path):
    """ดึงข้อความทุกหน้าของ PDF ออกมาเป็น list ของบรรทัด"""
    lines = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text(layout=False) or ''
            lines.extend(text.split('\n'))
    return lines


def parse_rows(lines):
    """แปลงบรรทัดข้อความให้เป็น list ของ dict ตามคอลัมน์ของรายงาน"""
    rows = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        # ข้ามหัวเรื่อง / ตัวกรองรายงาน / หัวตารางที่ขึ้นซ้ำทุกหน้า / เลขหน้า / แถวสรุปยอด
        if line.startswith('รายงานข้อมูลการใช้ไฟฟ้า'):
            continue
        if line.startswith('หน้าที่') and 'จาก' in line:
            continue
        if 'Number' in line and 'หน่วยงาน' in line:
            continue
        if line.startswith(('หน่วยงาน ', 'ชั้น ', 'โซน ', 'หมวด ')) and line.endswith('ทั้งหมด'):
            continue
        if line.startswith('วันที่'):
            continue
        if 'หน่วยใช้งาน' in line or line == 'ทั้งหมด':
            continue

        tokens = line.split()
        if len(tokens) < 9:
            continue

        try:
            number, onpeak, offpeak, total = (int(t) for t in tokens[-4:])
        except ValueError:
            continue

        rest = tokens[:-4]
        if len(rest) < 5 or rest[2] != 'Zone':
            continue

        rows.append({
            'หน่วยงาน': rest[0],
            'ชั้น': rest[1],
            'โซน': rest[2] + ' ' + rest[3],
            'บล็อก': rest[4],
            'หมวด': ' '.join(rest[5:]),
            'Number': number,
            'On-Peak Unit': onpeak,
            'Off-Peak Unit': offpeak,
            'Total Unit': total,
        })
    return rows


# ---------------------------------------------------------------------------
# 4) สร้างไฟล์ Excel
# ---------------------------------------------------------------------------
def build_excel(rows, output_path, sheet_title='ข้อมูลการใช้ไฟฟ้า'):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws.append(COLUMNS)
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER

    for r in rows:
        ws.append([r[col] for col in COLUMNS])

    last_row = ws.max_row
    for row_cells in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=len(COLUMNS)):
        for cell in row_cells:
            cell.font = NORMAL_FONT
            cell.border = BORDER
            if cell.column >= 6:
                cell.number_format = NUMFMT
                cell.alignment = Alignment(horizontal='right')
            else:
                cell.alignment = Alignment(horizontal='center' if cell.column != 5 else 'left')

    widths = [12, 8, 12, 10, 32, 10, 16, 16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = 'A2'
    table_name = re.sub(r'\W+', '_', f'ตาราง_{sheet_title}')
    tab = Table(displayName=table_name, ref=f'A1:{get_column_letter(len(COLUMNS))}{last_row}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tab)

    wb.save(output_path)
    return last_row - 1


# ---------------------------------------------------------------------------
# 5) ส่วนควบคุมหลัก: รองรับลากวางไฟล์ / ดับเบิลคลิก / รันผ่าน cmd
# ---------------------------------------------------------------------------
def resolve_output_path(pdf_path, explicit_output=None):
    if explicit_output:
        return explicit_output
    return os.path.splitext(pdf_path)[0] + '.xlsx'


def main():
    args = sys.argv[1:]
    pdf_args = [a for a in args if a.lower().endswith('.pdf')]

    explicit_output = None
    if len(pdf_args) == 1 and len(args) == 2 and args[1].lower().endswith('.xlsx'):
        explicit_output = args[1]

    pdf_files = [a for a in pdf_args if os.path.isfile(a)]

    if not pdf_files:
        # ไม่มีไฟล์ถูกลากมา หรือดับเบิลคลิกตรงๆ -> ค้นหาไฟล์ PDF ในโฟลเดอร์เดียวกับสคริปต์
        folder = os.path.dirname(os.path.abspath(__file__))
        pdf_files = sorted(glob.glob(os.path.join(folder, '*.pdf')))

    if not pdf_files:
        print('ไม่พบไฟล์ PDF')
        print('วิธีใช้: ลากไฟล์ PDF มาวางบนไอคอนนี้ หรือวางไฟล์ PDF ไว้ในโฟลเดอร์เดียวกันแล้วดับเบิลคลิกสคริปต์นี้')
        input('\nกด Enter เพื่อปิดหน้าต่างนี้...')
        return

    print(f'พบไฟล์ PDF จำนวน {len(pdf_files)} ไฟล์\n')

    success = 0
    for pdf_path in pdf_files:
        print(f'=== {os.path.basename(pdf_path)} ===')
        try:
            output_path = resolve_output_path(pdf_path, explicit_output if len(pdf_files) == 1 else None)
            lines = extract_lines(pdf_path)
            rows = parse_rows(lines)
            if not rows:
                print('  ไม่พบข้อมูลตารางในไฟล์นี้ (รูปแบบรายงานอาจไม่ตรงกัน) ข้ามไป\n')
                continue
            build_excel(rows, output_path)
            print(f'  สำเร็จ: พบ {len(rows)} แถว -> {os.path.basename(output_path)}\n')
            success += 1
        except Exception as e:
            print(f'  เกิดข้อผิดพลาด: {e}\n')

    print(f'เสร็จสิ้น: แปลงสำเร็จ {success}/{len(pdf_files)} ไฟล์')
    input('\nกด Enter เพื่อปิดหน้าต่างนี้...')


if __name__ == '__main__':
    main()
