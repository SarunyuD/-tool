# -*- coding: utf-8 -*-
"""
update_excel.py
================
ดึงข้อมูล "เต้ารับ" จากไฟล์ Excel ที่แปลงจาก PDF 
มาอัปเดตลงในไฟล์หลัก "การใช้พลังงานตามหมวดหมู่อุปกรณ์ ให้พี่เอก.xlsx"
และบันทึกเป็นไฟล์ใหม่ชื่อ "..._อัปเดตแล้ว.xlsx"
"""

import sys
import os
import subprocess
import importlib
import glob

# ---------------------------------------------------------------------------
# 1) ตรวจสอบและติดตั้งไลบรารี
# ---------------------------------------------------------------------------
def ensure_packages(packages):
    missing = []
    for pkg in packages:
        try:
            importlib.import_module(pkg)
        except ImportError:
            missing.append(pkg)
    if missing:
        print(f'กำลังติดตั้งไลบรารี ({", ".join(missing)})...')
        try:
            subprocess.check_call([sys.executable, '-m', 'pip', 'install', '--quiet', *missing])
            print('ติดตั้งสำเร็จ\n')
        except subprocess.CalledProcessError:
            print(f'ติดตั้งไม่สำเร็จ กรุณาพิมพ์ใน CMD: pip install {" ".join(missing)}')
            sys.exit(1)

ensure_packages(['openpyxl'])

import openpyxl

# ---------------------------------------------------------------------------
# 2) ค่าคงที่ชื่อไฟล์
# ---------------------------------------------------------------------------
MASTER_EXCEL_FILENAME = 'การใช้พลังงานตามหมวดหมู่อุปกรณ์ ให้พี่เอก.xlsx'

# ---------------------------------------------------------------------------
# 3) ฟังก์ชันอ่านข้อมูลจากไฟล์ Excel ที่แปลงแล้ว
# ---------------------------------------------------------------------------
def extract_data_from_source(source_path):
    print(f'กำลังอ่านข้อมูลจาก: {os.path.basename(source_path)} ...')
    wb = openpyxl.load_workbook(source_path, data_only=True)
    ws = wb.active
    
    # หาตำแหน่งคอลัมน์จาก Header ในแถวแรก
    headers = {str(cell.value).strip(): i+1 for i, cell in enumerate(ws[1]) if cell.value}
    
    col_agency = headers.get('หน่วยงาน')
    col_floor = headers.get('ชั้น')
    col_zone = headers.get('โซน')
    col_block = headers.get('บล็อก')
    col_category = headers.get('หมวด')
    
    col_number = headers.get('Number')
    col_onpeak = headers.get('On-Peak Unit')
    col_offpeak = headers.get('Off-Peak Unit')
    col_total = headers.get('Total Unit')

    if not all([col_agency, col_floor, col_zone, col_block, col_category]):
        print(f'  [ข้าม] โครงสร้างไฟล์ {os.path.basename(source_path)} ไม่ถูกต้อง (หาคอลัมน์ไม่ครบ)')
        return {}

    data_dict = {}
    for row in range(2, ws.max_row + 1):
        cat = str(ws.cell(row=row, column=col_category).value).strip() if ws.cell(row=row, column=col_category).value else ""
        
        if cat == 'เต้ารับ':
            a = str(ws.cell(row=row, column=col_agency).value).strip() if ws.cell(row=row, column=col_agency).value else ""
            f = str(ws.cell(row=row, column=col_floor).value).strip() if ws.cell(row=row, column=col_floor).value else ""
            z = str(ws.cell(row=row, column=col_zone).value).strip() if ws.cell(row=row, column=col_zone).value else ""
            b = str(ws.cell(row=row, column=col_block).value).strip() if ws.cell(row=row, column=col_block).value else ""
            
            key = (a, f, z, b)
            data_dict[key] = {
                'Number': ws.cell(row=row, column=col_number).value,
                'On-Peak Unit': ws.cell(row=row, column=col_onpeak).value,
                'Off-Peak Unit': ws.cell(row=row, column=col_offpeak).value,
                'Total Unit': ws.cell(row=row, column=col_total).value
            }
            
    print(f'  - พบข้อมูลหมวด "เต้ารับ" จำนวน {len(data_dict)} รายการ')
    return data_dict

# ---------------------------------------------------------------------------
# 4) ฟังก์ชันอัปเดตและบันทึกไฟล์หลัก (Save As)
# ---------------------------------------------------------------------------
def update_master_excel(master_path, all_data_dict):
    print(f'\n[กำลังเตรียมอัปเดตไฟล์หลัก] เปิดไฟล์ "{os.path.basename(master_path)}" ...')
    try:
        wb = openpyxl.load_workbook(master_path)
        ws = wb.active 

        update_count = 0
        for row in range(3, ws.max_row + 1):
            val_ao = ws[f'AO{row}'].value
            category = str(val_ao).strip() if val_ao is not None else ""
            
            if category == 'เต้ารับ':
                val_ak = ws[f'AK{row}'].value
                val_al = ws[f'AL{row}'].value
                val_am = ws[f'AM{row}'].value
                val_an = ws[f'AN{row}'].value

                k_agency = str(val_ak).strip() if val_ak is not None else ""
                k_floor  = str(val_al).strip() if val_al is not None else ""
                k_zone   = str(val_am).strip() if val_am is not None else ""
                k_block  = str(val_an).strip() if val_an is not None else ""

                key = (k_agency, k_floor, k_zone, k_block)
                
                if key in all_data_dict:
                    data = all_data_dict[key]
                    ws[f'AP{row}'] = data['Number']
                    ws[f'AQ{row}'] = data['On-Peak Unit']
                    ws[f'AS{row}'] = data['Off-Peak Unit']
                    ws[f'AU{row}'] = data['Total Unit']
                    update_count += 1

        # ตั้งชื่อไฟล์ใหม่โดยเติมคำว่า "_อัปเดตแล้ว" ต่อท้าย
        file_dir, file_name = os.path.split(master_path)
        name, ext = os.path.splitext(file_name)
        new_master_path = os.path.join(file_dir, f"{name}_อัปเดตแล้ว{ext}")

        # บันทึกเป็นไฟล์ใหม่
        wb.save(new_master_path)
        print(f'>> สำเร็จ: บันทึกข้อมูลเป็นไฟล์ใหม่ชื่อ "{os.path.basename(new_master_path)}" <<')
        print(f'>> (หยอดข้อมูลหมวด "เต้ารับ" ลงไปทั้งหมด {update_count} รายการ) <<\n')
        
    except PermissionError:
        print(f'>> [ข้อผิดพลาด] โปรดปิดไฟล์ Excel ต้นฉบับหรือไฟล์ที่กำลังจะเซฟก่อน แล้วรันใหม่อีกครั้ง <<\n')
    except Exception as e:
        print(f'>> [ข้อผิดพลาด] เกิดปัญหาขณะอัปเดตไฟล์หลัก: {e} <<\n')

# ---------------------------------------------------------------------------
# 5) ส่วนควบคุมหลัก
# ---------------------------------------------------------------------------
def main():
    folder = os.path.dirname(os.path.abspath(__file__))
    master_path = os.path.join(folder, MASTER_EXCEL_FILENAME)
    
    args = sys.argv[1:]
    # กรองไม่ให้ดึงไฟล์ "การใช้พลังงาน..." ทั้งแบบต้นฉบับและแบบที่เคยอัปเดตแล้วมาอ่าน
    source_files = [a for a in args if a.lower().endswith('.xlsx') 
                    and "การใช้พลังงาน" not in os.path.basename(a)]
    
    if not source_files:
        all_xlsx = glob.glob(os.path.join(folder, '*.xlsx'))
        source_files = [f for f in all_xlsx if "การใช้พลังงาน" not in os.path.basename(f)]

    if not source_files:
        print('ไม่พบไฟล์ Excel สำหรับดึงข้อมูล')
        print('คำแนะนำ: วางไฟล์ Excel ที่แปลงแล้วไว้ในโฟลเดอร์เดียวกัน หรือลากไฟล์มาใส่สคริปต์')
        input('\nกด Enter เพื่อปิดหน้าต่างนี้...')
        return

    if not os.path.isfile(master_path):
        print(f'ไม่พบไฟล์หลัก: "{MASTER_EXCEL_FILENAME}" ในโฟลเดอร์นี้')
        print('โปรดนำไฟล์หลักมาวางไว้โฟลเดอร์เดียวกันก่อนครับ')
        input('\nกด Enter เพื่อปิดหน้าต่างนี้...')
        return

    print(f'พบไฟล์ Excel ต้นทาง {len(source_files)} ไฟล์\n')
    
    combined_data = {}
    for src in source_files:
        try:
            data = extract_data_from_source(src)
            combined_data.update(data)
        except Exception as e:
            print(f'เกิดข้อผิดพลาดในการอ่านไฟล์ {os.path.basename(src)}: {e}')
            
    if combined_data:
        update_master_excel(master_path, combined_data)
    else:
        print('\nไม่พบข้อมูลหมวด "เต้ารับ" จากไฟล์ต้นทางเลย')

    input('\nกด Enter เพื่อปิดหน้าต่างนี้...')

if __name__ == '__main__':
    main()